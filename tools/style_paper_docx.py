"""Apply readable working-paper typography to the Pandoc Word export."""

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

__all__ = ["main"]


def element(name: str, **attrs: str) -> OxmlElement:
    """Create a WordprocessingML element with namespaced attributes."""
    node = OxmlElement("w:" + name)
    for key, value in attrs.items():
        node.set(qn("w:" + key), value)
    return node


def main(path: Path) -> None:
    """Style the existing document while preserving Pandoc's native footnotes."""
    doc = Document(path)
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for style in doc.styles:
        if style.type == 1:
            style.font.color.rgb = RGBColor(0, 0, 0)
            color = style.element.find(".//" + qn("w:color"))
            if color is not None:
                for attr in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn("w:" + attr), None)
        if style.name.startswith("Heading"):
            style.font.name = "Calibri"
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(16)
            style.paragraph_format.space_after = Pt(7)
        if style.name in ("Footnote Text", "Footnote"):
            style.font.name = "Georgia"
            style.font.size = Pt(8.5)
            style.paragraph_format.line_spacing = 1.0
            style.paragraph_format.space_after = Pt(2)
    for name, size in (
        ("Title", 25),
        ("Heading 1", 18),
        ("Heading 2", 15),
        ("Heading 3", 12),
    ):
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = section.right_margin = Inches(0.85)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.8)
        section.header_distance = section.footer_distance = Inches(0.3)
        header = section.header.paragraphs[0]
        header.text = "CULTIVATED CHICKEN AND HALAL MARKET ACCESS"
        header.style = doc.styles["Header"]
        header.runs[0].font.size = Pt(8)
        header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Working paper · 2 October 2026  |  ")
        run.font.size = Pt(8)
        field = element("fldSimple", instr="PAGE")
        footer._p.append(field)
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith(("Appendix ", "Bibliography")):
            paragraph.paragraph_format.page_break_before = True
        if paragraph.text.startswith("Figure "):
            paragraph.paragraph_format.space_before = Pt(7)
            paragraph.paragraph_format.space_after = Pt(9)
            for run in paragraph.runs:
                run.font.size = Pt(9)
        if paragraph._p.findall(".//" + qn("w:drawing")):
            paragraph.paragraph_format.keep_with_next = True
    for table in doc.tables:
        table.autofit = False
        table._tbl.tblPr.append(element("tblLayout", type="fixed"))
        borders = element("tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            borders.append(element(edge, val="single", sz="4", color="D9D9D9"))
        table._tbl.tblPr.append(borders)
        n = len(table.columns)
        weights = {
            2: [0.28, 0.72],
            3: [0.18, 0.49, 0.33],
            4: [0.20, 0.15, 0.30, 0.35],
        }.get(n, [1 / n] * n)
        for column, weight in zip(table.columns, weights):
            column.width = Inches(6.8 * weight)
        for i, row in enumerate(table.rows):
            if i == 0:
                row._tr.get_or_add_trPr().append(element("tblHeader", val="true"))
            row._tr.get_or_add_trPr().append(element("cantSplit"))
            for j, cell in enumerate(row.cells):
                cell.width = Inches(6.8 * weights[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = element("tcMar")
                for edge in ("top", "left", "bottom", "right"):
                    margins.append(element(edge, w="90", type="dxa"))
                props.append(margins)
                if i == 0:
                    props.append(element("shd", fill="DCE8ED"))
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.02
                    paragraph.paragraph_format.keep_with_next = i == 0
                    for run in paragraph.runs:
                        run.font.name = "Calibri"
                        run.font.size = Pt(9)
                        run.font.color.rgb = RGBColor(0, 0, 0)
                        if i == 0:
                            run.bold = True
    doc.core_properties.title = "Cultivated Chicken and Halal Market Access"
    doc.core_properties.subject = (
        "AI-assisted working paper; human verification pending"
    )
    doc.core_properties.author = "Muhammad Umar Zafar"
    doc.core_properties.keywords = (
        "cultivated chicken; halal; source audit; working paper"
    )
    doc.save(path)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
